"""Build regime_classifier_intro.pptx — matches LTCM_Presentation navy theme."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

REPO = Path(__file__).resolve().parents[1]
SLIDES_DIR = REPO / "Notebooks" / "output" / "regime_classifier_slides"
OUT_PPT = REPO / "regime_classifier_intro.pptx"
TEMPLATE = REPO / "LTCM_Presentation.pptx"

# LTCM_Presentation.pptx navy slides (e.g. slide 9)
BG_NAVY = RGBColor(0x0F, 0x1B, 0x2D)
TITLE_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SUBTITLE_CYAN = RGBColor(0x22, 0xD3, 0xEE)
BODY_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BODY_MUTED = RGBColor(0xCB, 0xD5, 0xE1)
TABLE_HEADER = RGBColor(0x1E, 0x3A, 0x5F)
TABLE_ROW = RGBColor(0x15, 0x27, 0x40)
ACCENT_CYAN = RGBColor(0x22, 0xD3, 0xEE)

TABLE_ROWS = [
    ("2Y trend", "10Y trend", "Spread trend", "Regime"),
    ("−1", "−1", "+1", "Bull Steepener"),
    ("−1", "−1", "−1", "Bull Flattener"),
    ("+1", "+1", "+1", "Bear Steepener"),
    ("+1", "+1", "−1", "Bear Flattener"),
    ("0", "—", "—", "Unclassified"),
]


def _set_slide_background(slide, color: RGBColor = BG_NAVY) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _style_para(p, text: str, *, size: int, bold: bool = False, color: RGBColor = BODY_WHITE) -> None:
    p.clear()
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def _blank_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_background(slide)
    return slide


def _add_title(slide, title: str, subtitle: str = "") -> None:
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.28), Inches(12.3), Inches(0.75))
    tf = box.text_frame
    tf.word_wrap = True
    _style_para(tf.paragraphs[0], title, size=24, bold=True, color=TITLE_WHITE)
    if subtitle:
        _style_para(tf.add_paragraph(), subtitle, size=11, color=SUBTITLE_CYAN)


def _add_bullets(slide, items: list[str], left=0.5, top=1.0, width=5.75, height=3.85, size=10) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        _style_para(p, item, size=size, color=BODY_MUTED)
        p.space_after = Pt(5)


def _add_formulas(slide, top=3.55) -> None:
    lines = [
        "d^{(i)} = SMA_{21}(y^{(i)}) − SMA_{200}(y^{(i)}),   i ∈ {2Y, 10Y}",
        "u = sign(d) with τ = 10 bp (yields), 5 bp (spread);  k = 5-day persistence",
    ]
    box = slide.shapes.add_textbox(Inches(0.5), Inches(top), Inches(5.75), Inches(0.55))
    tf = box.text_frame
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        _style_para(p, line, size=8.5, color=BODY_MUTED)


def _add_table(slide, top=4.15) -> None:
    rows, cols = len(TABLE_ROWS), 4
    shape = slide.shapes.add_table(rows, cols, Inches(0.5), Inches(top), Inches(5.75), Inches(1.35))
    table = shape.table
    for i, row in enumerate(TABLE_ROWS):
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = TABLE_HEADER if i == 0 else TABLE_ROW
            para = cell.text_frame.paragraphs[0]
            para.alignment = PP_ALIGN.CENTER if i == 0 else PP_ALIGN.LEFT
            _style_para(para, val, size=8 if i else 8.5, bold=(i == 0), color=TITLE_WHITE)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE


def _add_callout(slide, text: str, top=5.55) -> None:
    box = slide.shapes.add_textbox(Inches(0.5), Inches(top), Inches(12.3), Inches(0.5))
    p = box.text_frame.paragraphs[0]
    p.clear()
    r0 = p.add_run()
    r0.text = "Takeaway: "
    r0.font.size = Pt(9.5)
    r0.font.bold = True
    r0.font.color.rgb = ACCENT_CYAN
    r1 = p.add_run()
    r1.text = text
    r1.font.size = Pt(9.5)
    r1.font.color.rgb = BODY_WHITE


def _add_image(slide, name: str, left, top, width, height) -> None:
    path = SLIDES_DIR / name
    if path.exists():
        slide.shapes.add_picture(str(path), Inches(left), Inches(top),
                               width=Inches(width), height=Inches(height))


def main():
    import runpy
    runpy.run_path(str(Path(__file__).with_name("build_regime_slide_assets.py")), run_name="__main__")

    snap = pd.read_excel(REPO / "data" / "interest_rate_regimes.xlsx",
                         sheet_name="Current Snapshot").iloc[0]
    regime, as_of = snap["regime_filtered"], snap["as_of"]
    block_days = int(snap["block_days"])

    prs = Presentation()
    if TEMPLATE.exists():
        t = Presentation(str(TEMPLATE))
        prs.slide_width, prs.slide_height = t.slide_width, t.slide_height

    s1 = _blank_slide(prs)
    _add_title(s1, "Yield-Curve Regime Classifier",
               "Rule-based 2Y/10Y labels for portfolio conditioning")
    _add_bullets(s1, [
        "We map daily 2Y/10Y trends into six curve regimes plus Unclassified — "
        "a transparent overlay for when the commodity-trend sleeve behaves differently, not a trading signal.",
        "Mechanics: 21/200 SMA differentials, ±10 bp dead-zone on yields (±5 bp on spread), "
        "then a 5-day persistence filter so labels do not flip on one-day noise.",
        "Evidence: 2019+ chart tags COVID easing, 2022 hikes, higher-for-longer, and today’s bear-flatter block; "
        "regime-mix bar (2001-05-04+ panel) includes Unclassified (~52%) plus all six named regimes.",
    ], top=0.95, height=2.45)
    _add_formulas(s1)
    _add_table(s1)
    _add_image(s1, "slide1_timeseries_regimes.png", 6.15, 0.92, 6.55, 2.75)
    _add_image(s1, "slide1_named_regime_mix.png", 6.15, 3.78, 6.55, 1.45)
    _add_callout(s1, "Classifier recovers the same curve episodes macro investors already name.")

    s2 = _blank_slide(prs)
    _add_title(s2, "Robustness of Design Choices", "One-factor sensitivity on 2001+ sample")
    _add_bullets(s2, [
        "Method: vary one parameter at a time (k, τ_level, τ_spread, smoothing) while holding the rest at baseline — "
        "standard check before using labels in conditional performance work.",
        "Persistence k (1–10): Unclassified share stays ~51.5% (chart, left); median regime block grows 22 → 36 days. "
        "So ~half Unclassified is not a persistence artifact.",
        "τ_level (5–20 bp): main lever — 5 bp labels 63% of days; 20 bp leaves 70% Unclassified (chart, right). "
        "Spread τ and SMA 21/200 are stable vs alternatives; shorter windows over-label.",
        "Verdict: keep τ = 10/5 bp and k = 5 — conservative in chop (2021: ~94% Unclassified) "
        "but still captures 2022 (~5% Unclassified, sustained Bear Flattener).",
    ], top=0.95, height=4.35)
    _add_image(s2, "slide2_robustness_k_tau.png", 6.2, 1.0, 6.5, 3.65)
    _add_callout(s2, "Current thresholds are defensible; lowering τ trades false labels for coverage.")

    s3 = _blank_slide(prs)
    _add_title(s3, "Does the Label Fit Macro?", f"As of {as_of}: {regime} ({block_days}-day block)")
    _add_bullets(s3, [
        f"Today — {regime}: yields up, spread compressing (2Y {snap['y2_pct']}%, 10Y {snap['y10_pct']}%). "
        "Macro (May ’26): sticky inflation, elevated long-end/term premium, Fed on hold or hike repricing — "
        "bearish level move + flattening, not bull steepening from cuts.",
        "Forward (consensus, not our forecast): many desks flag bear steepener risk — long end on "
        "inflation/fiscal pressure, front end pinned; bull steepener needs material Fed cuts pulling 2Y down.",
        "History: 2020 Bull Steepener (COVID cuts); 2022 Bear Flattener (303-day block, ~5% Unclassified); "
        "2021 mostly Unclassified (low signal). Classifier confirms regime shifts — does not predict them.",
    ], top=0.95, height=4.35)
    _add_image(s3, "slide3_dominant_regime_by_year.png", 6.2, 1.0, 6.5, 3.65)
    _add_callout(s3,
                 "Today’s Bear Flattener fits the macro; if the curve steepens next, watch bear vs bull steepener — "
                 "our rules label it when 2Y/10Y trends confirm.")

    prs.save(str(OUT_PPT))
    print("Wrote", OUT_PPT)


if __name__ == "__main__":
    main()
