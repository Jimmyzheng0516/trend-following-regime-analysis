"""Build trend_performance_intro.pptx — matches LTCM_Presentation navy theme."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

REPO = Path(__file__).resolve().parents[1]
SLIDES_DIR = REPO / "Notebooks" / "output" / "trend_performance_slides"
OUT_PPT = REPO / "trend_performance_intro.pptx"
TEMPLATE = REPO / "LTCM_Presentation.pptx"

BG_NAVY = RGBColor(0x0F, 0x1B, 0x2D)
TITLE_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SUBTITLE_CYAN = RGBColor(0x22, 0xD3, 0xEE)
BODY_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BODY_MUTED = RGBColor(0xCB, 0xD5, 0xE1)
ACCENT_CYAN = RGBColor(0x22, 0xD3, 0xEE)


def _set_slide_background(slide, color: RGBColor = BG_NAVY) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _style_para(p, text: str, *, size: int, bold: bool = False, color: RGBColor = BODY_WHITE) -> None:
    p.clear()
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def _blank_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_background(slide)
    return slide


def _add_title(slide, title: str, subtitle: str = "") -> None:
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.28), Inches(12.3), Inches(0.75))
    tf = box.text_frame
    tf.word_wrap = True
    _style_para(tf.paragraphs[0], title, size=24, bold=True, color=TITLE_WHITE)
    if subtitle:
        _style_para(tf.add_paragraph(), subtitle, size=11, color=SUBTITLE_CYAN)


def _add_bullets(slide, items: list[str], left=0.5, top=1.0, width=5.75, height=4.35, size=10) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        _style_para(p, item, size=size, color=BODY_MUTED)
        p.space_after = Pt(5)


def _add_callout(slide, text: str, top=5.55) -> None:
    box = slide.shapes.add_textbox(Inches(0.5), Inches(top), Inches(12.3), Inches(0.5))
    p = box.text_frame.paragraphs[0]
    p.clear()
    r0 = p.add_run()
    r0.text = "Takeaway: "
    r0.font.size = Pt(9.5)
    r0.font.bold = True
    r0.font.color.rgb = ACCENT_CYAN
    r1 = p.add_run()
    r1.text = text
    r1.font.size = Pt(9.5)
    r1.font.color.rgb = BODY_WHITE


def _add_image(slide, name: str, left, top, width, height) -> None:
    path = SLIDES_DIR / name
    if path.exists():
        slide.shapes.add_picture(str(path), Inches(left), Inches(top),
                                 width=Inches(width), height=Inches(height))


def main() -> None:
    import runpy
    runpy.run_path(str(Path(__file__).with_name("build_trend_slide_assets.py")), run_name="__main__")

    prs = Presentation()
    if TEMPLATE.exists():
        t = Presentation(str(TEMPLATE))
        prs.slide_width, prs.slide_height = t.slide_width, t.slide_height

    s1 = _blank_slide(prs)
    _add_title(s1, "Commodity-Trend vs 60/40: Core Performance",
               "Method 1, $100M start, 2002–Apr 2026")
    _add_bullets(s1, [
        "Commodity-Trend is not a standalone replacement for 60/40: 4.38% annualized vs 7.67%, "
        "ending $283M vs $603M — but with 7.1% vol vs 10.9% and Sharpe 0.62 vs 0.71.",
        "Max drawdown -15.9% vs -32.7% (GFC trough); worst daily -3.9% vs -6.5% — "
        "a smoother path, not a higher-return engine.",
        "Framing for the team: evaluate as a sleeve inside 60/40, not as a product on its own.",
    ], top=0.95, height=4.0)
    _add_image(s1, "slide1_drawdown_episodes_timeline.png", 6.2, 0.92, 4.85, 2.5)
    _add_image(s1, "slide1_core_metrics_table.png", 6.2, 3.48, 4.85, 1.05)
    _add_image(s1, "slide1_annual_returns_stacked.png", 8.62, 0.92, 4.46, 3.68)
    _add_callout(s1, "Lower standalone return; materially lower tail risk — sets up why crisis windows matter.")

    s2 = _blank_slide(prs)
    _add_title(s2, "Crisis Drawdowns: Where the Sleeve Adds Value",
               "Peak-to-trough returns in the deepest 13 60/40 episodes")
    _add_bullets(s2, [
        "In the two largest 60/40 crises, Trend gained while 60/40 lost: "
        "GFC +4.9% vs -32.4%; 2022 +3.0% vs -19.2%.",
        "Across all 13 deepest episodes, Trend is flat-to-positive in 5 of 13 vs 60/40 losing in every case; "
        "max Trend episode loss -4.1% (2025) vs 60/40 -32.4% (GFC).",
        "Honest limits: COVID 2020 (-0.7% vs -19.3%) and 2025 (-4.1% vs -10.7%) — "
        "dampens losses, not a perfect hedge. In worst 10% of 63-day 60/40 windows, "
        "Trend +0.2% vs 60/40 -6.9% (+7.2pp).",
    ], top=0.95, height=4.35)
    _add_image(s2, "slide2_deepest13_drawdown_bars.png", 6.2, 0.95, 6.5, 3.85)
    _add_callout(s2,
                 "The sleeve earns its keep in large 60/40 drawdowns, especially GFC and 2022 — not every stress episode.")

    s3 = _blank_slide(prs)
    _add_title(s3, "When Does It Work? Conditional Performance",
               "Monthly win rates — vs sleeves and when stock-bond correlation is positive")
    _add_bullets(s3, [
        "vs sleeves: Trend beats equity only 38.1% of months outright, "
        "but 83.5% when equity is down (+3.5pp avg active). vs bonds: 50.5% outright, 72.9% when bonds are down (+1.1pp).",
        "Not a return replacement — value is conditional hedging when a traditional sleeve is losing.",
        "Positive stock-bond correlation (~25% of months): Trend beats 60/40 only 40.3% of the time in those months alone.",
        "Stress overlap is what matters: win rate rises to 69.6% when pos-corr months are also down for 60/40; "
        "77.8% when both stocks and bonds are down (+1.7pp).",
    ], top=0.95, height=4.35)
    _add_image(s3, "slide3_part5_sleeve_win_rates.png", 6.2, 0.95, 6.5, 1.95)
    _add_image(s3, "slide3_part6_poscorr_stress.png", 6.2, 2.98, 6.5, 1.95)
    _add_callout(s3,
                 "Trend hedges failing sleeves and adds most vs 60/40 when diversification breaks and the portfolio is already losing.")

    prs.save(str(OUT_PPT))
    print("Wrote", OUT_PPT)


if __name__ == "__main__":
    main()
