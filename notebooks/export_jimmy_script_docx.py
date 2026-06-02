"""Export jimmy_part_presenter_script.md structure to Word."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor

REPO = Path(__file__).resolve().parents[1]
MD = REPO / "Notebooks" / "output" / "jimmy_part_presenter_script.md"
OUT = REPO / "Notebooks" / "output" / "jimmy_part_presenter_script.docx"


def _add_bullet(doc: Document, text: str, level: int = 0) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.25 + 0.25 * level)
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.font.name = "Calibri"


def _add_heading(doc: Document, text: str, level: int) -> None:
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = "Calibri"


def _add_label_para(doc: Document, label: str, body: str, *, color: RGBColor | None = None) -> None:
    p = doc.add_paragraph()
    r0 = p.add_run(f"{label} ")
    r0.bold = True
    r0.font.size = Pt(11)
    r0.font.name = "Calibri"
    if color:
        r0.font.color.rgb = color
    r1 = p.add_run(body)
    r1.font.size = Pt(11)
    r1.font.name = "Calibri"
    p.paragraph_format.space_after = Pt(6)


def _add_meta(doc: Document, line: str) -> None:
    p = doc.add_paragraph(line.strip("* ").replace("**", ""))
    for run in p.runs:
        run.font.size = Pt(10)
        run.font.name = "Calibri"
        run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)


def _add_body(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text)
    for run in p.runs:
        run.font.size = Pt(11)
        run.font.name = "Calibri"


def main() -> None:
    lines = MD.read_text(encoding="utf-8").splitlines()
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    _add_heading(doc, "Jimmy's Part — Full Presenter Script", level=0)

    in_qa = False
    in_timing = False
    skip_next_empty_after_say = False

    i = 0
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()

        if not line:
            i += 1
            continue
        if line == "---":
            i += 1
            continue
        if line.startswith("# ") and "Presenter Script" in line:
            i += 1
            continue

        if line.startswith("**Deck:") or line.startswith("**Model:") or line.startswith("**Benchmark:") or line.startswith("**Total time:"):
            _add_meta(doc, line)
            i += 1
            continue

        if line.startswith("## "):
            title = line[3:].strip()
            in_qa = title.startswith("Q&A")
            in_timing = title.startswith("Timing cheat sheet")
            _add_heading(doc, title, level=1)
            i += 1
            continue

        if line.startswith("**Takeaway:**"):
            _add_label_para(doc, "Takeaway:", line.replace("**Takeaway:**", "").strip(),
                            color=RGBColor(0x22, 0xD3, 0xEE))
            i += 1
            continue

        if line in ("**Say:**", "**Close:**"):
            label = line.strip("*").replace(":", "")
            p = doc.add_paragraph()
            run = p.add_run(f"{label}:")
            run.bold = True
            run.italic = True
            run.font.size = Pt(11)
            run.font.name = "Calibri"
            i += 1
            continue

        if line.startswith("**Point to:**"):
            _add_label_para(doc, "Point to:", line.replace("**Point to:**", "").strip())
            i += 1
            continue

        if in_timing:
            if line.startswith("|") and not line.startswith("|-"):
                cells = [c.strip() for c in line.strip("|").split("|")]
                if cells and cells[0] != "Slide":
                    _add_bullet(doc, f"Slide {cells[0]} — {cells[1]} — pause: {cells[2]}", level=0)
            elif line.startswith("**If short:**"):
                _add_body(doc, line.replace("**", ""))
            i += 1
            continue

        if in_qa and line.strip().startswith("**") and line.strip().endswith("**"):
            q = line.strip("*").strip()
            i += 1
            answer_parts = []
            while i < len(lines):
                nxt = lines[i].rstrip()
                if not nxt:
                    i += 1
                    if answer_parts:
                        break
                    continue
                if nxt.strip().startswith("**") and nxt.strip().endswith("**"):
                    break
                if nxt.startswith("## "):
                    break
                answer_parts.append(nxt)
                i += 1
            p = doc.add_paragraph()
            rq = p.add_run(q)
            rq.bold = True
            rq.font.size = Pt(11)
            rq.font.name = "Calibri"
            p.paragraph_format.space_before = Pt(8)
            ans = " ".join(p.strip("- ").strip() for p in answer_parts)
            ans = re.sub(r"\*\*(.+?)\*\*", r"\1", ans)
            if ans:
                _add_body(doc, ans)
            continue

        m = re.match(r"^(\s*)- (.+)$", raw)
        if m:
            spaces = len(m.group(1))
            content = re.sub(r"\*\*(.+?)\*\*", r"\1", m.group(2))
            _add_bullet(doc, content, level=spaces // 2)
            i += 1
            continue

        # plain paragraph (e.g. Q&A answer on same line after question - handled above)
        if in_qa and line and not line.startswith("#"):
            _add_body(doc, line)
            i += 1
            continue

        i += 1

    doc.save(str(OUT))
    print("Wrote", OUT)


if __name__ == "__main__":
    main()
