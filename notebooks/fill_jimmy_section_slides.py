"""Fill section-divider slides in Jimmy's-part-LTCM-Presentation.pptx."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

REPO = Path(__file__).resolve().parents[1]
PPT = REPO / "Jimmy's-part-LTCM-Presentation.pptx"

# y-position bands (EMU) for the four text boxes on divider slides
Y_TITLE = 1_097_280
Y_QUESTION = 1_920_240
Y_PREVIEW = 3_054_600
X_PREVIEW_LEFT = 812_773
X_PREVIEW_RIGHT = 5_186_629

SECTIONS: dict[int, tuple[str, str, str, str]] = {
    1: (
        "Commodity-Trend vs 60/40",
        "Does a commodity-trend sleeve improve a traditional 60/40 portfolio — "
        "especially when stocks, bonds, or both are losing?",
        "Core performance: 4.4% ann. vs 7.7%, but −16% max drawdown vs −33%",
        "Crisis drawdowns + conditional hedging when equity or bond sleeves fail",
    ),
    5: (
        "Yield-Curve Regime Classifier",
        "How do we label daily 2Y/10Y curve regimes so we can explain "
        "when the trend sleeve behaves differently?",
        "Rule-based labels: SMA21−SMA200, ±10 bp dead-zones, 5-day persistence",
        "Robustness checks + macro validation (COVID, 2022 hikes, today)",
    ),
    10: (
        "Commodity-Trend × Yield-Curve Regimes",
        "When does commodity-trend add value — and what happens when the regime label changes?",
        "Steepening/twist → return edge; flatteners → diversification",
        "65 regime transitions since 2001: Bull Steepener vs Bear Flattener paths",
    ),
}


def _shape_pos(shape) -> tuple[int | None, int | None]:
    if not hasattr(shape, "left"):
        return None, None
    return int(shape.left), int(shape.top)


def _set_shape_text(shape, text: str) -> None:
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text


def _find_shape_by_y(slide, y_target: int, x_target: int | None = None):
    best = None
    best_dist = float("inf")
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        left, top = _shape_pos(shape)
        if top is None:
            continue
        if abs(top - y_target) > 50_000:
            continue
        if x_target is not None and left is not None:
            dist = abs(top - y_target) + abs(left - x_target)
        else:
            dist = abs(top - y_target)
        if dist < best_dist:
            best_dist = dist
            best = shape
    return best


def fill_section_slide(slide, title: str, question: str, preview_left: str, preview_right: str) -> None:
    mapping = [
        (Y_TITLE, None, title),
        (Y_QUESTION, None, question),
        (Y_PREVIEW, X_PREVIEW_LEFT, preview_left),
        (Y_PREVIEW, X_PREVIEW_RIGHT, preview_right),
    ]
    for y, x, text in mapping:
        shape = _find_shape_by_y(slide, y, x)
        if shape is None:
            raise RuntimeError(f"Could not find text shape at y={y}, x={x}")
        _set_shape_text(shape, text)


def main() -> None:
    prs = Presentation(str(PPT))
    for slide_idx, content in SECTIONS.items():
        slide = prs.slides[slide_idx - 1]
        fill_section_slide(slide, *content)
        print(f"Filled slide {slide_idx}: {content[0]}")

    prs.save(str(PPT))
    print("Wrote", PPT)


if __name__ == "__main__":
    main()
